"""
Document service.

Responsibilities:
- Extract text from PDF / DOCX / TXT files.
- Persist document metadata (filename, size, status, chunk count, errors)
  in a local SQLite database — separate from Qdrant, which only stores
  vectors + chunk payloads.
- Provide listing / status-update / deletion of document records.

Ingestion (chunking + embedding + Qdrant upsert) is intentionally kept in
rag_service.py so extraction and retrieval logic stay separated per the
project's architecture requirement.
"""

import logging
import sqlite3
from contextlib import contextmanager
from dataclasses import dataclass
from pathlib import Path
from typing import List, Optional

from pypdf import PdfReader
from docx import Document as DocxDocument

from config.settings import settings
from utils.helpers import now_iso

logger = logging.getLogger(__name__)

STATUS_UPLOADING = "uploading"
STATUS_EXTRACTING = "extracting"
STATUS_CHUNKING = "chunking"
STATUS_EMBEDDING = "embedding"
STATUS_INDEXING = "indexing"
STATUS_COMPLETED = "completed"
STATUS_FAILED = "failed"


class ExtractionError(Exception):
    def __init__(self, code: str, message: str):
        self.code = code
        self.message = message
        super().__init__(message)


@dataclass
class ExtractedDocument:
    pages: List[str]  # one entry per page; for txt/docx this is a single "page"


def _init_db():
    settings.DATABASE_PATH.parent.mkdir(parents=True, exist_ok=True)
    with sqlite3.connect(settings.DATABASE_PATH) as conn:
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS documents (
                id TEXT PRIMARY KEY,
                filename TEXT NOT NULL,
                file_type TEXT NOT NULL,
                file_size INTEGER NOT NULL,
                upload_date TEXT NOT NULL,
                status TEXT NOT NULL,
                chunk_count INTEGER NOT NULL DEFAULT 0,
                error_message TEXT
            )
            """
        )
        conn.commit()


@contextmanager
def _get_conn():
    conn = sqlite3.connect(settings.DATABASE_PATH)
    conn.row_factory = sqlite3.Row
    try:
        yield conn
    finally:
        conn.close()


def create_document_record(doc_id: str, filename: str, file_type: str, file_size: int) -> None:
    with _get_conn() as conn:
        conn.execute(
            """INSERT INTO documents (id, filename, file_type, file_size, upload_date, status, chunk_count)
               VALUES (?, ?, ?, ?, ?, ?, 0)""",
            (doc_id, filename, file_type, file_size, now_iso(), STATUS_UPLOADING),
        )
        conn.commit()


def update_status(doc_id: str, status: str, error_message: Optional[str] = None) -> None:
    with _get_conn() as conn:
        conn.execute(
            "UPDATE documents SET status = ?, error_message = ? WHERE id = ?",
            (status, error_message, doc_id),
        )
        conn.commit()


def set_chunk_count(doc_id: str, chunk_count: int) -> None:
    with _get_conn() as conn:
        conn.execute("UPDATE documents SET chunk_count = ? WHERE id = ?", (chunk_count, doc_id))
        conn.commit()


def list_documents() -> List[dict]:
    with _get_conn() as conn:
        rows = conn.execute("SELECT * FROM documents ORDER BY upload_date DESC").fetchall()
        return [dict(r) for r in rows]


def get_document(doc_id: str) -> Optional[dict]:
    with _get_conn() as conn:
        row = conn.execute("SELECT * FROM documents WHERE id = ?", (doc_id,)).fetchone()
        return dict(row) if row else None


def delete_document_record(doc_id: str) -> None:
    with _get_conn() as conn:
        conn.execute("DELETE FROM documents WHERE id = ?", (doc_id,))
        conn.commit()


def knowledge_base_stats() -> dict:
    with _get_conn() as conn:
        row = conn.execute(
            "SELECT COUNT(*) AS docs, COALESCE(SUM(chunk_count), 0) AS chunks "
            "FROM documents WHERE status = ?",
            (STATUS_COMPLETED,),
        ).fetchone()
        return {"documents": row["docs"], "chunks": row["chunks"]}


# --- Text extraction ---------------------------------------------------

def extract_text(file_path: Path, extension: str) -> ExtractedDocument:
    """Extract text from a supported document. Raises ExtractionError on failure."""
    try:
        if extension == ".pdf":
            return _extract_pdf(file_path)
        if extension == ".docx":
            return _extract_docx(file_path)
        if extension == ".txt":
            return _extract_txt(file_path)
    except ExtractionError:
        raise
    except Exception as exc:  # noqa: BLE001 - convert any parser failure into a clean error
        logger.exception("Text extraction failed for %s", file_path)
        raise ExtractionError("EXTRACTION_FAILED", f"Could not read this file: {exc}") from exc

    raise ExtractionError("UNSUPPORTED_TYPE", f"Unsupported file extension: {extension}")


def _extract_pdf(file_path: Path) -> ExtractedDocument:
    try:
        reader = PdfReader(str(file_path))
    except Exception as exc:  # noqa: BLE001
        raise ExtractionError("CORRUPTED_FILE", f"This PDF could not be opened: {exc}") from exc

    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception:  # noqa: BLE001
            raise ExtractionError("ENCRYPTED_PDF", "This PDF is password-protected and cannot be read.")

    pages = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception:  # noqa: BLE001
            pages.append("")

    if not any(p.strip() for p in pages):
        raise ExtractionError(
            "NO_TEXT_FOUND",
            "No extractable text was found in this PDF (it may be a scanned image without OCR).",
        )

    return ExtractedDocument(pages=pages)


def _extract_docx(file_path: Path) -> ExtractedDocument:
    try:
        doc = DocxDocument(str(file_path))
    except Exception as exc:  # noqa: BLE001
        raise ExtractionError("CORRUPTED_FILE", f"This DOCX could not be opened: {exc}") from exc

    paragraphs = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
    table_text = []
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                table_text.append(" | ".join(cells))

    full_text = "\n\n".join(paragraphs + table_text)
    if not full_text.strip():
        raise ExtractionError("NO_TEXT_FOUND", "No extractable text was found in this DOCX file.")

    return ExtractedDocument(pages=[full_text])


def _extract_txt(file_path: Path) -> ExtractedDocument:
    raw = file_path.read_bytes()
    for encoding in ("utf-8", "utf-8-sig", "latin-1"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ExtractionError("ENCODING_ERROR", "This TXT file's encoding could not be determined.")

    if not text.strip():
        raise ExtractionError("NO_TEXT_FOUND", "This TXT file is empty of usable text.")

    return ExtractedDocument(pages=[text])


_init_db()
